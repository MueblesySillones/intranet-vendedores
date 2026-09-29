# -*- coding: utf-8 -*-
"""De donde salen las filas.

Al resto del panel no le importa si la planilla es un CSV que alguien exporto
del sistema, un Excel que se manda por mail o —mas adelante— una hoja de
Google. Todos piden lo mismo: `filas`, una lista de listas donde `filas[0]`
son los encabezados. Este modulo es el unico que sabe abrir archivos.

    leer(config) -> {"ok", "filas", "origen", "cuando", "desde_cache", ...}

`config` describe la fuente: {"tipo": "csv"|"xlsx", "ruta": ..., "hoja": ...}.
Si falta `tipo` se deduce de la extension.

QUE HACE ADEMAS DE ABRIR EL ARCHIVO
  · Adivina la codificacion. El Excel en castellano exporta en cp1252 o en
    utf-8 con BOM segun el dia; leerlo mal no rompe nada, solo deja
    "MaÃ±ana" en el nombre de una sucursal y ese nombre despues no agrupa.
  · Adivina el separador. "Guardar como CSV" en una Windows en castellano
    escribe con ";" porque la coma es el separador decimal.
  · Saltea las filas vacias y los titulos que suele haber arriba del
    encabezado real.
  · Desambigua columnas con el mismo nombre.
  · Guarda el resultado en un cache en disco para no releer 50.000 filas en
    cada click.

QUE NO HACE, A PROPOSITO
  · No rellena las filas cortas. Una fila a la que le faltan celdas es un
    error de carga que `revisor.py` detecta comparando `len(fila)` contra
    `len(encabezado)`. Emparejarlas aca dejaria la tabla prolija y muda.
  · No inventa nombres para las columnas sin titulo: `analizador.py` las
    ignora a proposito, asi que solo avisa que estan.
  · No calcula ni corrige nada. Si algo no se puede leer con certeza lo dice
    en `error` o en `avisos`, que es la regla 5 del contrato.

Solo biblioteca estandar. `openpyxl` es opcional y se importa con try/except:
sin el, los .xlsx no se pueden leer pero el panel arranca igual.
"""
import codecs
import csv
import datetime
import hashlib
import io
import json
import os
import re
import sys
from collections import Counter

try:
    import openpyxl
except ImportError:          # el panel tiene que arrancar igual sin openpyxl
    openpyxl = None


# ── topes ────────────────────────────────────────────────────────────────
# El panel es un servidor local de un solo proceso: una lectura que tarda un
# minuto lo deja mudo para todos. Preferimos cortar y decirlo.
MAX_BYTES = 60 * 1024 * 1024      # ~600.000 filas de CSV; mas que eso no es una planilla
MAX_FILAS = 200000                # la de derivaciones tiene 1.800; el peor caso real ronda 50.000
MUESTRA = 262144                  # lo que se lee para adivinar codificacion y separador

# Sube el tope del modulo csv (131.072 por defecto). Una celda con un texto
# largo pegado adentro hacia fallar la planilla ENTERA con "field larger than
# field limit", que no le dice nada a nadie. No se usa sys.maxsize: en Windows
# desborda el long de C y tira OverflowError.
csv.field_size_limit(10 * 1024 * 1024)

EXT_CSV = {".csv", ".txt", ".tsv"}
EXT_XLSX = {".xlsx", ".xlsm"}
SEPARADORES = [",", ";", "\t", "|"]

# Si cambia como se parsea, los caches viejos quedan mintiendo. Subir este
# numero los invalida a todos de una.
VERSION_CACHE = 1
CACHE_MAX_ARCHIVOS = 40


class _Corte(Exception):
    """Un problema previsto, con un mensaje escrito para la persona que lo ve.

    Se usa en vez de devolver tuplas (datos, error) por todos lados: el error
    viaja solo hasta `leer()`, que lo convierte en {"ok": False, "error": ...}.
    """


# ── donde vive el cache ──────────────────────────────────────────────────
def carpeta_estado():
    """La carpeta de estado del panel, SIEMPRE fuera del arbol del proyecto.

    Importa mas de lo que parece: el cache guarda las filas tal cual, o sea
    nombres, telefonos y mails de clientes. Si cayera adentro del repo, un
    `git add` distraido lo publicaria en la intranet, que no tiene contrasena.
    Por eso en modo desarrollo NO se imita `dirname(EXE_DIR)` como hace el
    panel (ahi eso da `herramientas/`, adentro del proyecto) y se va a
    %LOCALAPPDATA%.
    """
    forzada = os.environ.get("MYS_PANEL_STATE")
    if forzada:
        return forzada
    if getattr(sys, "frozen", False):
        # empaquetado: el mismo PanelMyS_state que usa panel_server.py
        return os.path.join(os.path.dirname(os.path.dirname(sys.executable)),
                            "PanelMyS_state")
    base = os.environ.get("LOCALAPPDATA") or os.path.expanduser("~")
    return os.path.join(base, "PanelMyS_state")


def carpeta_cache():
    return os.path.join(carpeta_estado(), "cache_planillas")


# ── utilidades chicas ────────────────────────────────────────────────────
def _valor(v):
    """Una celda, siempre como texto.

    Que CSV y XLSX devuelvan exactamente lo mismo no es cosmetica: asi el
    cache puede ser JSON sin perder nada, y `analizador.py` ve la misma tabla
    venga de donde venga.
    """
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):
        # Excel guarda las fechas como instantes. Si es medianoche era una
        # fecha; dejarla como "2026-08-13 00:00:00" hace que el rango
        # desde/hasta del analizador, que ordena textos, mezcle formatos.
        if (v.hour, v.minute, v.second) == (0, 0, 0):
            return v.strftime("%Y-%m-%d")
        return v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, datetime.date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, datetime.time):
        return v.strftime("%H:%M")
    if isinstance(v, float):
        # Excel guarda TODO numero como float: "12 personas" llega como 12.0 y
        # una categoria armada con eso separaria "12" de "12.0".
        if v == int(v) and abs(v) < 1e15:
            return str(int(v))
        return str(v)
    return str(v)


def _no_vacias(fila):
    return sum(1 for v in fila if str(v).strip())


def _ahora():
    return datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")


# ── codificacion ─────────────────────────────────────────────────────────
def _codificacion(muestra, avisos):
    """Con que codificacion hay que abrir este CSV.

    El orden no es caprichoso: el BOM es una firma, no una adivinanza, asi que
    manda. Despues se prueba utf-8 ESTRICTO, que se valida a si mismo —un
    archivo cp1252 con acentos casi nunca forma secuencias utf-8 validas—, y
    recien ahi se cae a cp1252, que es lo que escribe el Excel en Windows.
    Al reves (probar cp1252 primero) nunca falla, porque cp1252 acepta casi
    cualquier byte, y todo archivo utf-8 quedaria mal leido en silencio.
    """
    if muestra.startswith(codecs.BOM_UTF8):
        return "utf-8-sig"
    # UTF-32LE empieza con los mismos dos bytes que UTF-16LE, por eso va antes
    if muestra.startswith(codecs.BOM_UTF32_LE) or muestra.startswith(codecs.BOM_UTF32_BE):
        return "utf-32"
    if muestra.startswith(codecs.BOM_UTF16_LE) or muestra.startswith(codecs.BOM_UTF16_BE):
        return "utf-16"     # el "Texto Unicode" del Excel
    # final=False: la muestra corta el archivo en cualquier lado y puede partir
    # un caracter de varios bytes al medio. Sin esto, un utf-8 grande se
    # declaraba cp1252 por culpa del ultimo caracter, que estaba cortado.
    try:
        codecs.getincrementaldecoder("utf-8")().decode(muestra, False)
        return "utf-8"
    except UnicodeDecodeError:
        pass
    try:
        muestra.decode("cp1252")
        return "cp1252"
    except UnicodeDecodeError:
        avisos.append("El archivo tiene bytes que no son ni utf-8 ni cp1252. Se leyo "
                      "igual, pero puede haber caracteres cambiados: conviene abrirlo "
                      "en Excel y volver a guardarlo como CSV UTF-8.")
        return "latin-1"    # no falla nunca: los 256 bytes tienen un caracter


# ── separador ────────────────────────────────────────────────────────────
def _separador(texto, cortado, avisos):
    """Con que caracter estan separadas las columnas.

    No se usa csv.Sniffer: falla con archivos de una sola columna y elige
    cualquier cosa cuando hay comillas. Aca se prueba cada candidato de
    verdad, con el propio csv.reader, y gana el que parte TODAS las filas en
    la misma cantidad de columnas. Esa consistencia es la senal que de verdad
    distingue un separador de una coma decimal.
    """
    lineas = [l for l in texto.splitlines()[:40] if l.strip()]
    # la ultima linea de una muestra recortada esta partida al medio y contaria
    # como una fila inconsistente que en el archivo no existe
    if cortado and len(lineas) > 1:
        lineas = lineas[:-1]
    if not lineas:
        return ","

    mejor, puntaje = ",", (0.0, 0)
    for cand in SEPARADORES:
        try:
            filas = [f for f in csv.reader(lineas, delimiter=cand) if _no_vacias(f)]
        except csv.Error:
            continue
        if not filas:
            continue
        anchos = Counter(len(f) for f in filas)
        ancho, veces = anchos.most_common(1)[0]
        if ancho < 2:
            continue                     # con este caracter no separo nada
        p = (veces / float(len(filas)), ancho)
        if p > puntaje:
            mejor, puntaje = cand, p

    if puntaje[0] and puntaje[0] < 0.9:
        avisos.append("Separando por %s las filas no dan todas la misma cantidad de "
                      "columnas. Fijate que la planilla no tenga celdas corridas."
                      % _nombre_sep(mejor))
    return mejor


def _nombre_sep(s):
    return {"\t": "tabulador", ";": "punto y coma", ",": "coma", "|": "barra"}.get(s, s)


# ── acomodar la tabla ────────────────────────────────────────────────────
def _sacar_arriba(filas, avisos, mirar=25):
    """Tira lo que haya arriba del encabezado real.

    Casi toda planilla que sale de un sistema trae un titulo, un logo o dos
    filas en blanco antes de los nombres de columna. Si no se sacan, el
    analizador toma el titulo como encabezado y la planilla entera queda con
    una sola columna llamada "Reporte de derivaciones".
    """
    if not filas:
        return filas
    i = 0
    while i < len(filas) and _no_vacias(filas[i]) == 0:
        i += 1                     # una fila 100% vacia nunca es el encabezado
    filas = filas[i:]
    fuera = i
    if not filas:
        return filas

    cuentas = [_no_vacias(f) for f in filas[:mirar]]
    ancho = max(cuentas)
    # Con una sola columna cualquier umbral se comeria la tabla entera.
    if ancho > 1:
        # Mitad del ancho, y nunca menos de 2. El encabezado real puede tener
        # una columna sin nombre; un titulo tiene UNA sola celda escrita. El
        # umbral tiene que dejar pasar al primero y frenar al segundo.
        minimo = max(2, int(ancho * 0.5))
        j = 0
        while j < len(cuentas) - 1 and j < 10 and cuentas[j] < minimo:
            j += 1
        filas = filas[j:]
        fuera += j
    if fuera:
        avisos.append("Se saltearon %d fila(s) de arriba (vacias o de titulo) hasta "
                      "encontrar el encabezado: %s"
                      % (fuera, ", ".join(str(c) for c in filas[0][:4] if str(c).strip())))
    return filas


def _sacar_abajo(filas, avisos):
    """Tira las filas vacias del FINAL, no las del medio.

    Excel marca como usadas las filas donde alguien apoyo el cursor, y un xlsx
    con 1.800 datos puede declarar 30.000 filas. Las del medio se dejan a
    proposito: son un error de carga real y `revisor.py` las reporta.
    """
    n = len(filas)
    while filas and _no_vacias(filas[-1]) == 0:
        filas.pop()
    if n - len(filas):
        avisos.append("Se ignoraron %d fila(s) vacias del final de la planilla."
                      % (n - len(filas)))
    return filas


def _sacar_columnas_fantasma(filas, avisos):
    """Recorta las columnas de la derecha que no tienen NADA, ni siquiera titulo."""
    if not filas:
        return filas
    ancho = max(len(f) for f in filas)
    tope = 0
    for f in filas:
        j = len(f)
        while j > tope and not str(f[j - 1]).strip():
            j -= 1
        if j > tope:
            tope = j
            if tope >= ancho:
                return filas       # no hay nada para recortar; cortamos temprano
    if tope and tope < ancho:
        avisos.append("Se ignoraron %d columna(s) del final que estaban vacias."
                      % (ancho - tope))
        return [f[:tope] for f in filas]
    return filas


def _encabezados(cab, cuerpo, avisos):
    """Normaliza los titulos y desambigua los repetidos.

    Los repetidos importan porque `revisor.py` arma un diccionario
    {nombre: columna}: con dos columnas "Seguimiento" una se pierde, y los
    avisos de la que sobrevive se leen como si fueran de la otra. Renombrar la
    segunda a "Seguimiento (2)" cuesta una linea y evita un aviso grave
    apuntando a la columna equivocada.
    """
    limpios, vistos, repetidos, sin_nombre = [], {}, [], []
    for i, c in enumerate(cab):
        # el alt+enter del Excel mete saltos de linea adentro del titulo, y
        # " Sucursal " y "Sucursal" se veian como dos columnas distintas
        n = re.sub(r"\s+", " ", str(c).replace("\n", " ").replace("\t", " ")).strip()
        if not n:
            # No se le inventa un nombre: `analizador.py` ignora a proposito
            # las columnas sin titulo. Solo se avisa, y solo si tienen datos.
            limpios.append("")
            if any(str(f[i]).strip() for f in cuerpo if i < len(f)):
                sin_nombre.append(i + 1)
            continue
        k = n.lower()
        if k in vistos:
            vistos[k] += 1
            repetidos.append(n)
            n = "%s (%d)" % (n, vistos[k])
        else:
            vistos[k] = 1
        limpios.append(n)

    if repetidos:
        avisos.append("Hay columnas con el mismo titulo (%s). Se les agrego un numero "
                      "para poder distinguirlas." % ", ".join(sorted(set(repetidos))))
    if sin_nombre:
        avisos.append("La(s) columna(s) %s tienen datos pero no tienen titulo, asi que "
                      "no se van a analizar. Ponele un nombre en la planilla."
                      % ", ".join(str(x) for x in sin_nombre))
    return limpios


def _acomodar(filas, avisos):
    filas = _sacar_arriba(filas, avisos)
    filas = _sacar_abajo(filas, avisos)
    filas = _sacar_columnas_fantasma(filas, avisos)
    if filas:
        filas[0] = _encabezados(filas[0], filas[1:], avisos)
    return filas


# ── abrir el archivo ─────────────────────────────────────────────────────
def _revisar_archivo(ruta):
    """Que el archivo exista, se pueda abrir y no sea absurdamente grande.

    Se chequea ANTES de parsear nada para que el error hable de lo que pasa de
    verdad —el archivo esta abierto en Excel, la ruta cambio— y no de un
    traceback de csv o de zipfile tres capas mas abajo.
    """
    if not ruta or not str(ruta).strip():
        raise _Corte("No dijiste que archivo abrir.")
    ruta = os.path.abspath(os.path.expanduser(str(ruta).strip().strip('"')))

    if os.path.basename(ruta).startswith("~$"):
        # Excel deja un ~$archivo.xlsx al lado mientras lo tenes abierto; es un
        # archivo de bloqueo de 160 bytes, no la planilla, y aparece en el
        # selector de archivos igual que la buena.
        raise _Corte("Ese no es el archivo: los que empiezan con ~$ son los archivos "
                     "temporales que crea Excel mientras la planilla esta abierta. "
                     "Elegi el que tiene el nombre sin el ~$.")
    if os.path.isdir(ruta):
        raise _Corte("«%s» es una carpeta, no una planilla." % os.path.basename(ruta))
    if not os.path.exists(ruta):
        raise _Corte("No encontre el archivo:\n%s\nFijate que no lo hayan movido, "
                     "renombrado, o que no este en una carpeta compartida que hoy no "
                     "esta conectada." % ruta)
    try:
        tam = os.path.getsize(ruta)
    except OSError as e:
        raise _Corte("No pude ni mirar el tamano del archivo (%s)." % e)
    if tam == 0:
        raise _Corte("El archivo esta vacio (0 bytes). Si lo acabas de exportar, puede "
                     "ser que el sistema todavia lo este escribiendo.")
    if tam > MAX_BYTES:
        raise _Corte("El archivo pesa %.0f MB y el tope es %d MB. Una planilla asi no "
                     "se puede abrir sin dejar el panel colgado: partila por periodo "
                     "y cargala en dos veces."
                     % (tam / 1048576.0, MAX_BYTES // 1048576))
    # Abrirlo y cerrarlo es la unica forma barata de saber si Windows lo deja.
    try:
        open(ruta, "rb").close()
    except PermissionError:
        raise _Corte("Windows no me deja leer el archivo. Casi siempre es porque esta "
                     "abierto en Excel: cerralo del todo y probá de nuevo.")
    except OSError as e:
        raise _Corte("No se pudo abrir el archivo: %s" % e)
    return ruta, tam


def _tipo_de(ruta, cfg):
    tipo = (cfg.get("tipo") or "").strip().lower()
    if tipo in ("csv", "xlsx"):
        return tipo
    ext = os.path.splitext(ruta)[1].lower()
    if ext in EXT_CSV:
        return "csv"
    if ext in EXT_XLSX:
        return "xlsx"
    if ext == ".xls":
        # openpyxl no lee el binario viejo de Excel 97 y no hay nada en la
        # biblioteca estandar que lo haga.
        raise _Corte("Los .xls (Excel viejo) no se pueden leer. Abrilo en Excel y "
                     "guardalo como .xlsx o como CSV.")
    if ext in (".ods", ".numbers", ".gsheet"):
        raise _Corte("El formato %s no se puede leer. Exportalo como CSV o .xlsx." % ext)
    raise _Corte("No se que tipo de archivo es «%s». Se pueden leer CSV (.csv, .txt, "
                 ".tsv) y Excel (.xlsx, .xlsm)." % (ext or os.path.basename(ruta)))


def _leer_csv(ruta, cfg, avisos):
    with open(ruta, "rb") as f:
        muestra = f.read(MUESTRA)
    cortado = os.path.getsize(ruta) > MUESTRA

    enc = (cfg.get("codificacion") or "").strip() or _codificacion(muestra, avisos)
    # El separador se busca sobre el TEXTO ya decodificado, no sobre los bytes:
    # en un archivo utf-16 cada caracter lleva un \x00 al lado y contar bytes
    # da cualquier cosa.
    # Un byte NUL no existe en un CSV de texto. Hasta Python 3.10 el modulo csv
    # cortaba solo con "line contains NUL"; desde la 3.11 los deja pasar, asi
    # que un .db o una foto renombrada a .csv se leia "bien" y devolvia una
    # unica fila de basura. Se chequea a mano. UTF-16/32 quedan afuera porque
    # ahi el NUL es parte normal de cada caracter.
    if b"\x00" in muestra and not enc.startswith(("utf-16", "utf-32")):
        raise _Corte("El archivo tiene bytes binarios adentro: no es un CSV de texto. "
                     "Si le cambiaste la extension a otro archivo, o lo guardaste como "
                     "«Texto Unicode», exportalo de nuevo como CSV UTF-8.")

    texto = muestra.decode(enc, "replace")
    sep = cfg.get("separador") or _separador(texto, cortado, avisos)

    tope = int(cfg.get("max_filas") or MAX_FILAS)
    filas, truncado = [], False
    try:
        with io.open(ruta, encoding=enc, newline="") as f:   # newline="" lo pide el modulo csv
            for fila in csv.reader(f, delimiter=sep):
                if len(filas) >= tope:
                    truncado = True
                    break
                filas.append([_valor(v) for v in fila])
    except UnicodeDecodeError:
        # La muestra decia una cosa y mas abajo aparecio un byte que no cierra.
        # Se relee entero en latin-1, que no puede fallar, y se avisa fuerte:
        # es mejor una planilla con dos caracteres raros que ninguna planilla.
        avisos.append("A partir de cierta fila el archivo dejo de ser %s. Se releyo "
                      "todo en latin-1: revisa que los acentos y las enes se vean "
                      "bien antes de publicar cualquier numero." % enc)
        enc = "latin-1"
        filas, truncado = [], False
        with io.open(ruta, encoding=enc, newline="") as f:
            for fila in csv.reader(f, delimiter=sep):
                if len(filas) >= tope:
                    truncado = True
                    break
                filas.append([_valor(v) for v in fila])
    except csv.Error as e:
        if "NUL" in str(e):
            raise _Corte("El archivo tiene bytes binarios adentro: no es un CSV de "
                         "verdad. Si lo renombraste a .csv, exportalo bien desde el "
                         "programa que lo genero.")
        raise _Corte("El CSV esta mal armado y no se pudo terminar de leer: %s" % e)

    return filas, truncado, {"codificacion": enc, "separador": sep, "hoja": "", "hojas": []}


def _hoja_elegida(wb, pedida):
    nombres = wb.sheetnames
    if pedida in (None, "", -1):
        # No se usa wb.active: es la hoja que estaba seleccionada cuando
        # guardaron, o sea cualquiera. La primera VISIBLE es predecible, y las
        # ocultas suelen ser tablas auxiliares que nadie quiere ver.
        for n in nombres:
            try:
                if wb[n].sheet_state == "visible":
                    return n
            except (KeyError, AttributeError):
                pass
        return nombres[0]
    if isinstance(pedida, int) or str(pedida).strip().lstrip("-").isdigit():
        i = int(pedida)
        if not (0 <= i < len(nombres)):
            raise _Corte("El archivo tiene %d hoja(s) y pediste la numero %d."
                         % (len(nombres), i))
        return nombres[i]
    # sin distinguir mayusculas ni espacios: el nombre va y vuelve por la URL
    # del panel y no siempre vuelve identico
    quiere = str(pedida).strip().lower()
    for n in nombres:
        if n.strip().lower() == quiere:
            return n
    raise _Corte("La hoja «%s» no existe en el archivo. Las que hay son: %s."
                 % (pedida, ", ".join(nombres)))


def _leer_xlsx(ruta, cfg, avisos):
    if openpyxl is None:
        raise _Corte("Para leer archivos de Excel hace falta la libreria openpyxl y "
                     "en esta computadora no esta instalada. Mientras tanto, abri la "
                     "planilla y guardala como CSV (Archivo > Guardar como > CSV UTF-8): "
                     "eso se lee igual de bien.")
    tope = int(cfg.get("max_filas") or MAX_FILAS)
    wb = None
    try:
        try:
            # read_only: no carga la hoja entera en memoria, va fila por fila.
            # data_only: devuelve el VALOR que Excel dejo guardado, no la
            # formula. Sin esto una columna calculada llega como "=B2*C2".
            wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
        except PermissionError:
            raise _Corte("Windows no me deja leer el archivo. Casi siempre es porque "
                         "esta abierto en Excel: cerralo del todo y probá de nuevo.")
        except Exception as e:      # noqa — openpyxl tira de todo: BadZipFile, KeyError, ...
            if e.__class__.__name__ == "BadZipFile":
                raise _Corte("Ese archivo no es un Excel de verdad, aunque termine en "
                             ".xlsx. Suele pasar cuando se le cambia la extension a un "
                             "CSV: dejalo con su extension original.")
            raise _Corte("No se pudo abrir el Excel: %s" % e)

        hojas_todas = list(wb.sheetnames)
        hoja = _hoja_elegida(wb, cfg.get("hoja"))
        ws = wb[hoja]

        filas, truncado = [], False
        for fila in ws.iter_rows(values_only=True):     # values_only ahorra crear una celda por dato
            if len(filas) >= tope:
                truncado = True
                break
            filas.append([_valor(v) for v in fila])
    finally:
        # Sin close(), en read_only el zip queda abierto y Windows no deja
        # sobrescribir el archivo: el usuario no puede guardar desde Excel
        # hasta cerrar el panel.
        if wb is not None:
            try:
                wb.close()
            except Exception:       # noqa
                pass

    if len(hojas_todas) > 1 and cfg.get("hoja") in (None, ""):
        avisos.append("El archivo tiene %d hojas y se leyo «%s». Las otras son: %s."
                      % (len(hojas_todas), hoja,
                         ", ".join(n for n in hojas_todas if n != hoja)))
    return filas, truncado, {"codificacion": "", "separador": "",
                             "hoja": hoja, "hojas": hojas_todas}


def _formulas_sin_calcular(ruta, hoja, filas):
    """Nombres de columna que quedaron vacias porque tienen formulas sin valor.

    `data_only=True` devuelve el ultimo valor que CALCULO Excel. Si la planilla
    la genero un programa, o si la guardaron sin recalcular, ese valor no
    existe y la columna llega vacia sin ninguna senal. Un dashboard con una
    columna calculada en cero es exactamente el numero que se ve bien y esta
    mal, asi que hay que decirlo.

    Cuesta abrir el archivo una segunda vez, por eso solo se llama cuando ya
    sabemos que hay alguna columna 100% vacia, y solo mira las primeras filas.
    """
    if openpyxl is None or len(filas) < 2:
        return []
    cab = filas[0]
    sospechosas = [i for i, n in enumerate(cab)
                   if str(n).strip() and
                   not any(str(f[i]).strip() for f in filas[1:] if i < len(f))]
    if not sospechosas:
        return []
    wb = None
    con_formula = []
    try:
        wb = openpyxl.load_workbook(ruta, read_only=True, data_only=False)
        ws = wb[hoja]
        for n, fila in enumerate(ws.iter_rows(values_only=True)):
            if n > 200:
                break
            for i in list(sospechosas):
                v = fila[i] if i < len(fila) else None
                if isinstance(v, str) and v.startswith("="):
                    con_formula.append(cab[i])
                    sospechosas.remove(i)
            if not sospechosas:
                break
    except Exception:       # noqa — es un chequeo de cortesia; si falla, no pasa nada
        return []
    finally:
        if wb is not None:
            try:
                wb.close()
            except Exception:   # noqa
                pass
    return con_formula


# ── cache en disco ───────────────────────────────────────────────────────
def _marca(ruta):
    """La huella del archivo: si esta cambia, el cache no sirve mas.

    Tamano + fecha de modificacion alcanzan el 99% de las veces, pero fallan
    justo en el caso que mas pasa: alguien corrige una celda y guarda dentro
    del mismo segundo sin cambiar el largo. Por eso se agregan 64 KB del
    principio y 64 KB del final, que cuestan milisegundos aunque la planilla
    pese 30 MB.
    """
    st = os.stat(ruta)
    h = hashlib.sha1()
    with open(ruta, "rb") as f:
        h.update(f.read(65536))
        if st.st_size > 131072:
            f.seek(-65536, os.SEEK_END)
            h.update(f.read(65536))
    return {
        # redondeado a milisegundos: en carpetas de red el mtime vuelve con
        # precision distinta segun quien lo pregunte, y comparando el float
        # crudo el cache no servia nunca
        "mtime": round(st.st_mtime, 3),
        "tam": st.st_size,
        "bordes": h.hexdigest(),
        "version": VERSION_CACHE,
    }


def _clave(ruta, cfg):
    """(clave del archivo, clave de como se lo lee).

    Van separadas para poder borrar todo lo cacheado de UN archivo sin abrir
    ningun cache. `max_filas` entra en la segunda: si no, una lectura recortada
    para la vista previa se serviria despues como si fuera la planilla entera.
    """
    a = hashlib.sha1(os.path.normcase(ruta).encode("utf-8")).hexdigest()[:12]
    resto = "|".join(str(cfg.get(k, "")) for k in
                     ("hoja", "separador", "codificacion", "max_filas"))
    return a, hashlib.sha1(resto.encode("utf-8")).hexdigest()[:6]


def _ruta_cache(ruta, cfg):
    a, b = _clave(ruta, cfg)
    return os.path.join(carpeta_cache(), "planilla_%s_%s.json" % (a, b))


def _cache_leer(ruta, cfg, marca):
    p = _ruta_cache(ruta, cfg)
    if not os.path.isfile(p):
        return None
    try:
        with io.open(p, encoding="utf-8") as f:
            guardado = json.load(f)
    except (ValueError, OSError):
        # Un cache a medio escribir (panel cerrado de golpe) envenenaria todas
        # las lecturas siguientes. Se borra y se lee el archivo de nuevo.
        try:
            os.remove(p)
        except OSError:
            pass
        return None
    if guardado.get("marca") != marca:
        return None
    res = guardado.get("res")
    return res if isinstance(res, dict) and res.get("filas") is not None else None


def _cache_guardar(ruta, cfg, marca, res):
    p = _ruta_cache(ruta, cfg)
    try:
        os.makedirs(os.path.dirname(p), exist_ok=True)
        # Se escribe en un temporal y se renombra: os.replace es atomico, asi
        # que dos pestanas del panel leyendo a la vez no pueden dejar un JSON
        # cortado por la mitad.
        tmp = "%s.%d.tmp" % (p, os.getpid())
        with io.open(tmp, "w", encoding="utf-8") as f:
            json.dump({"marca": marca, "res": res}, f, ensure_ascii=False)
        os.replace(tmp, p)
        _podar()
    except OSError:
        pass        # sin cache el panel anda igual, solo mas lento


def _podar():
    """El cache no puede crecer para siempre en el disco de otra persona."""
    try:
        d = carpeta_cache()
        archivos = [os.path.join(d, n) for n in os.listdir(d) if n.endswith(".json")]
        if len(archivos) <= CACHE_MAX_ARCHIVOS:
            return
        archivos.sort(key=lambda p: os.path.getmtime(p))
        for p in archivos[:len(archivos) - CACHE_MAX_ARCHIVOS]:
            try:
                os.remove(p)
            except OSError:
                pass
    except OSError:
        pass


def limpiar_cache(ruta=None):
    """Borra el cache de un archivo, o todo. Es el boton "volver a leer"."""
    d = carpeta_cache()
    if not os.path.isdir(d):
        return 0
    prefijo = "planilla_"
    if ruta:
        prefijo += _clave(os.path.abspath(str(ruta)), {})[0]
    n = 0
    for nombre in os.listdir(d):
        if nombre.startswith(prefijo):
            try:
                os.remove(os.path.join(d, nombre))
                n += 1
            except OSError:
                pass
    return n


# ── lo que usa el resto del panel ────────────────────────────────────────
def hojas(ruta):
    """Los nombres de las hojas, para poder elegir ANTES de leer la planilla.

    Un .xlsx con cinco hojas no se puede leer a ciegas, y leerlo entero solo
    para llenar un desplegable es tirar diez segundos a la basura.
    """
    try:
        ruta, _ = _revisar_archivo(ruta)
        if _tipo_de(ruta, {}) != "xlsx":
            return {"ok": True, "hojas": []}       # un CSV es una sola hoja
        if openpyxl is None:
            raise _Corte("Para mirar las hojas de un Excel hace falta openpyxl y no "
                         "esta instalado.")
        wb = None
        try:
            wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
            return {"ok": True, "hojas": list(wb.sheetnames)}
        finally:
            if wb is not None:
                try:
                    wb.close()
                except Exception:   # noqa
                    pass
    except _Corte as e:
        return {"ok": False, "error": str(e), "hojas": []}
    except Exception as e:          # noqa
        return {"ok": False, "error": "No se pudo mirar el archivo: %s" % e, "hojas": []}


def leer(config):
    """Las filas de una planilla, venga de donde venga.

    config = {"ruta": ..., "tipo": "csv"|"xlsx" (opcional, sale de la
    extension), "hoja": nombre o numero (xlsx), "separador", "codificacion"
    (opcional, se adivinan), "max_filas" (opcional), "cache": False para
    forzar la relectura}

    Devuelve SIEMPRE un dict, nunca levanta una excepcion por un problema
    previsto: el panel tiene que poder mostrar el motivo, y "el archivo esta
    abierto en Excel" no es un bug, es el martes.

      ok           True/False
      filas        [[encabezados...], [fila], ...]  (vacia si ok es False)
      origen       de donde salio, para poder rastrear un numero despues
      cuando       cuando se leyo el ARCHIVO (no cuando se sirvio el cache)
      desde_cache  si no hizo falta volver a abrirlo
      avisos       cosas que la persona tiene que saber (lista de textos)
      truncado     True si la planilla era mas larga que el tope: los totales
                   estan incompletos y NO se pueden publicar
      error        solo cuando ok es False
    """
    cfg = dict(config or {})
    avisos = []
    try:
        ruta, tam = _revisar_archivo(cfg.get("ruta"))
        tipo = _tipo_de(ruta, cfg)
        usar_cache = cfg.get("cache", True) is not False
        marca = _marca(ruta)

        if usar_cache:
            guardado = _cache_leer(ruta, cfg, marca)
            if guardado is not None:
                guardado = dict(guardado)
                guardado["desde_cache"] = True
                return guardado

        if tipo == "csv":
            filas, truncado, extra = _leer_csv(ruta, cfg, avisos)
        else:
            filas, truncado, extra = _leer_xlsx(ruta, cfg, avisos)

        filas = _acomodar(filas, avisos)
        if not filas:
            raise _Corte("El archivo se abrio bien pero no tiene ninguna fila con "
                         "datos adentro.")
        if truncado:
            # Primero en la lista y con mayusculas a proposito: una planilla
            # cortada da totales reales pero incompletos, que es la peor clase
            # de numero equivocado porque no se nota.
            avisos.insert(0, "ATENCION: la planilla tiene mas de %d filas y se leyeron "
                             "solo las primeras. Todo lo que se cuente abajo esta "
                             "INCOMPLETO y no se puede publicar."
                          % int(cfg.get("max_filas") or MAX_FILAS))
        if len(filas) == 1:
            avisos.append("Solo aparece el encabezado: la planilla no tiene filas de datos.")

        if tipo == "xlsx":
            con_formula = _formulas_sin_calcular(ruta, extra["hoja"], filas)
            if con_formula:
                avisos.append("La(s) columna(s) %s tienen formulas que Excel nunca "
                              "calculo, asi que llegan vacias. Abri la planilla en "
                              "Excel, guardala y volve a cargarla."
                              % ", ".join(con_formula))

        res = {
            "ok": True,
            "filas": filas,
            "origen": ruta + (" · hoja %s" % extra["hoja"] if extra["hoja"] else ""),
            "cuando": _ahora(),
            "desde_cache": False,
            "ruta": ruta,
            "archivo": os.path.basename(ruta),
            "tipo": tipo,
            "hoja": extra["hoja"],
            "hojas": extra["hojas"],
            "codificacion": extra["codificacion"],
            "separador": extra["separador"],
            "peso": tam,
            "total_filas": len(filas) - 1,
            "total_columnas": len(filas[0]),
            "truncado": truncado,
            "avisos": avisos,
        }
        if usar_cache:
            _cache_guardar(ruta, cfg, marca, res)
        return res

    except _Corte as e:
        return {"ok": False, "error": str(e), "filas": [], "origen": str(cfg.get("ruta") or ""),
                "cuando": _ahora(), "desde_cache": False, "avisos": avisos, "truncado": False}
    except MemoryError:
        return {"ok": False, "error": "La planilla no entra en memoria. Partila por "
                                      "periodo y cargala en dos veces.",
                "filas": [], "origen": str(cfg.get("ruta") or ""), "cuando": _ahora(),
                "desde_cache": False, "avisos": avisos, "truncado": False}
    except Exception as e:      # noqa — el panel nunca puede quedarse sin respuesta
        return {"ok": False,
                "error": "No se pudo leer la planilla (%s: %s)" % (e.__class__.__name__, e),
                "filas": [], "origen": str(cfg.get("ruta") or ""), "cuando": _ahora(),
                "desde_cache": False, "avisos": avisos, "truncado": False}


if __name__ == "__main__":
    import sys as _sys
    _sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    aca = os.path.dirname(os.path.abspath(__file__))
    ruta = _sys.argv[1] if len(_sys.argv) > 1 else os.path.join(
        os.path.dirname(aca), "derivaciones_prueba.csv")
    r = leer({"ruta": ruta, "hoja": _sys.argv[2] if len(_sys.argv) > 2 else ""})
    if not r["ok"]:
        print("NO SE PUDO LEER\n%s" % r["error"])
        _sys.exit(1)
    print("%s  ·  %d filas x %d columnas%s"
          % (r["archivo"], r["total_filas"], r["total_columnas"],
             "  (del cache)" if r["desde_cache"] else ""))
    if r["codificacion"]:
        print("leido como %s, separado por %s" % (r["codificacion"], _nombre_sep(r["separador"])))
    if r["hoja"]:
        print("hoja: %s   (hay %d)" % (r["hoja"], len(r["hojas"])))
    print("cuando: %s" % r["cuando"])
    for a in r["avisos"]:
        print("  ! %s" % a)
    print("\nencabezados: %s" % ", ".join(str(c) for c in r["filas"][0]))
    for f in r["filas"][1:4]:
        print("   %s" % " | ".join(str(v)[:18] for v in f))
